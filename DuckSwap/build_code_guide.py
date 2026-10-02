from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Users\vince\OneDrive\Documents\ChatGPT\Game development\Guia tecnica del codigo Patitos en batalla.docx")
DUCK = Path(r"C:\Users\vince\OneDrive\Documents\ChatGPT\Game development\DuckSwap\DuckPreview.png")
BLACK = RGBColor(0,0,0)
NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F3F6FA"
GRAY = "D9D9D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color=GRAY):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tb = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),color); tb.append(e)
    tblPr.append(tb)

def margins(cell, top=90, start=110, bottom=90, end=110):
    tc=cell._tc; pr=tc.get_or_add_tcPr(); m=pr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); pr.append(m)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        x=m.find(qn('w:'+name))
        if x is None: x=OxmlElement('w:'+name); m.append(x)
        x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa')

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9.5)
    set_repeat_header(t.rows[0])
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=str(v); margins(cs[i]); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cs[i],PALE)
            for p in cs[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(9.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    borders(t); doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def code(doc, text, caption=None):
    if caption:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
        r=p.add_run(caption); r.bold=True; r.font.size=Pt(9.5)
        spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(1); spacer.paragraph_format.line_spacing=Pt(1)
        spacer.add_run(' ').font.size=Pt(1)
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.6)
    c=t.cell(0,0); shade(c,"F5F5F5"); margins(c,130,160,130,160)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text.strip()); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(35,35,35)
    borders(t,"C7C7C7"); doc.add_paragraph().paragraph_format.space_after=Pt(0)

def bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.add_run(text); return p

def source(doc, path, lines):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(f"Fuente del proyecto  {path}  líneas {lines}"); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(90,90,90)

def page(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=BLACK
styles['Normal'].paragraph_format.space_after=Pt(7); styles['Normal'].paragraph_format.line_spacing=1.12
styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(30); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=BLACK
title_ppr=styles['Title']._element.get_or_add_pPr()
old_border=title_ppr.find(qn('w:pBdr'))
if old_border is not None: title_ppr.remove(old_border)
no_border=OxmlElement('w:pBdr')
for edge in ('top','left','bottom','right','between','bar'):
    node=OxmlElement('w:'+edge); node.set(qn('w:val'),'nil'); no_border.append(node)
title_ppr.append(no_border)
for name,size in [('Heading 1',20),('Heading 2',14),('Heading 3',11.5)]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=BLACK; s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(12 if name!='Heading 1' else 8); s.paragraph_format.space_after=Pt(6)
if 'Code Caption' not in styles:
    styles.add_style('Code Caption', WD_STYLE_TYPE.PARAGRAPH)

# Footer
for section in doc.sections:
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.add_run('Guía técnica del código  Patitos en batalla   •   ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
    for r in f.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

# Cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(45); p.add_run('Guía técnica del código')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22)
r=p.add_run('Patitos en batalla'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(23); r.font.color.rgb=BLACK
if DUCK.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(DUCK),width=Inches(3.1))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18)
r=p.add_run('Documento complementario del GDD'); r.bold=True; r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Unity 2019.4.19f1  •  C Sharp  •  Escena principal Main').font.size=Pt(10)
page(doc)

doc.add_heading('1 Propósito y alcance',level=1)
doc.add_paragraph('Este documento explica las partes del código que sostienen la versión modificada del juego. Su propósito es ayudar a presentar el proyecto, reconocer qué script controla cada comportamiento y facilitar cambios posteriores sin tener que leer todo el repositorio.')
doc.add_paragraph('La conclusión técnica principal es que el personaje visual cambió de tanque a pato de goma, pero la arquitectura conserva los nombres originales Tank. El prefab sigue reuniendo movimiento, disparo, salud, física y UI; el modelo de pato funciona como su representación visual.')
doc.add_heading('Mapa de responsabilidades',level=2)
table(doc,['Script','Responsabilidad principal'],[
('GameManager.cs','Inicia la partida, crea jugadores y administra rondas y ganadores.'),
('GameTimerPause.cs','Cuenta el tiempo, actualiza la interfaz y pausa o reanuda el juego.'),
('TankManager.cs','Conecta cada personaje con su número, color, aparición y controles.'),
('TankMovement.cs','Lee ejes de entrada y mueve o gira el Rigidbody.'),
('TankShooting.cs','Carga el disparo, crea el proyectil y aplica su velocidad.'),
('TankHealth.cs','Gestiona vida, barra de salud, muerte y explosión.'),
('ShellExplosion.cs','Detecta impactos, calcula daño por distancia y aplica fuerza.')],[1.75,4.75])
doc.add_heading('Flujo general',level=2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Main.unity  →  GameManager.Start  →  SpawnAllTanks  →  GameLoop  →  ronda  →  ganador'); r.bold=True; r.font.size=Pt(11)
page(doc)

doc.add_heading('2 Inicio de la partida',level=1)
doc.add_paragraph('No existe actualmente un botón de inicio independiente. Unity carga la escena Main y llama automáticamente a Start en GameManager. Ese método prepara las esperas, crea los personajes, registra los objetivos de cámara y comienza la corrutina que gobierna la partida.')
code(doc,"""private void Start()
{
    m_StartWait = new WaitForSeconds(m_StartDelay);
    m_EndWait = new WaitForSeconds(m_EndDelay);

    SpawnAllTanks();
    SetCameraTargets();
    StartCoroutine(GameLoop());
}""",'Inicio automático en GameManager')
source(doc,'Assets/Scripts/Managers/GameManager.cs','24 a 34')
doc.add_heading('Creación de los personajes',level=2)
code(doc,"""for (int i = 0; i < m_Tanks.Length; i++)
{
    m_Tanks[i].m_Instance = Instantiate(
        m_TankPrefab,
        m_Tanks[i].m_SpawnPoint.position,
        m_Tanks[i].m_SpawnPoint.rotation);
    m_Tanks[i].m_PlayerNumber = i + 1;
    m_Tanks[i].Setup();
}""",'SpawnAllTanks crea una instancia por jugador')
doc.add_paragraph('Cada elemento de m_Tanks contiene un punto de aparición y un color. El mismo prefab se reutiliza para ambos jugadores; el número asignado decide qué entradas del teclado leerá cada instancia.')
doc.add_heading('Qué sería un botón de inicio',level=2)
doc.add_paragraph('Si se desea agregarlo en el futuro, el menú debe mantener Time.timeScale en cero o impedir el inicio de GameLoop hasta que el botón invoque un método público. Esa función no forma parte de la versión actual y conviene implementarla en un controlador de menú separado para no mezclarla con la lógica de rondas.')
page(doc)

doc.add_heading('3 Temporizador y pausa',level=1)
doc.add_paragraph('GameTimerPause es la mejora principal de interfaz. Awake restablece la escala temporal, enlaza el botón con TogglePause y dibuja el valor inicial. Update escucha Escape y aumenta el tiempo solo cuando el juego no está pausado.')
code(doc,"""private void Awake()
{
    Time.timeScale = 1f;
    if (pauseButton != null)
        pauseButton.onClick.AddListener(TogglePause);
    UpdateDisplay();
}

private void Update()
{
    if (Input.GetKeyDown(KeyCode.Escape))
        TogglePause();

    if (!isPaused)
    {
        elapsedTime += Time.unscaledDeltaTime;
        UpdateDisplay();
    }
}""",'Preparación y actualización del HUD')
source(doc,'Assets/Scripts/Managers/GameTimerPause.cs','13 a 30')
doc.add_heading('Cómo se detiene el juego',level=2)
code(doc,"""public void TogglePause()
{
    isPaused = !isPaused;
    Time.timeScale = isPaused ? 0f : 1f;
    if (pauseButtonText != null)
        pauseButtonText.text = isPaused ? "REANUDAR" : "PAUSAR";
}""",'Un mismo método atiende al botón y a Escape')
doc.add_paragraph('Time.timeScale controla la velocidad del tiempo simulado. Al usar 0, la física y los métodos dependientes de deltaTime se detienen. El texto del botón comunica el siguiente estado disponible.')
page(doc)

doc.add_heading('4 Lectura y presentación del tiempo',level=1)
code(doc,"""private void UpdateDisplay()
{
    if (timerText == null) return;
    int totalSeconds = Mathf.FloorToInt(elapsedTime);
    timerText.text = string.Format(
        "TIEMPO  {0:00}:{1:00}",
        totalSeconds / 60,
        totalSeconds % 60);
}""",'Conversión de segundos al formato minutos y segundos')
source(doc,'Assets/Scripts/Managers/GameTimerPause.cs','41 a 46')
doc.add_paragraph('elapsedTime almacena segundos como float. FloorToInt descarta la fracción para mostrar segundos completos. La división entera calcula minutos y el módulo obtiene los segundos restantes. El formato 00 asegura dos dígitos, por ejemplo 02:07.')
doc.add_heading('Por qué usa unscaledDeltaTime',level=2)
doc.add_paragraph('Time.unscaledDeltaTime mide el tiempo real entre cuadros sin multiplicarlo por Time.timeScale. En este proyecto el contador solo se incrementa dentro de if (!isPaused), de modo que se detiene durante la pausa por decisión explícita. Esta combinación evita depender accidentalmente del estado global del motor.')
doc.add_heading('Protecciones incluidas',level=2)
bullet(doc,'Las referencias timerText y pauseButtonText se validan antes de usarse, lo que evita errores NullReferenceException si falta un enlace en el Inspector.')
bullet(doc,'OnDestroy restablece Time.timeScale a 1 para que otra escena o una nueva ejecución no quede congelada.')
bullet(doc,'El mismo TogglePause se reutiliza para el clic y la tecla Escape, por lo que ambos caminos producen exactamente el mismo estado.')
page(doc)

doc.add_heading('5 Controles de los jugadores',level=1)
table(doc,['Acción','Jugador 1','Jugador 2','Entrada lógica'],[
('Avanzar y retroceder','W y S','Flechas arriba y abajo','Vertical1 o Vertical2'),
('Girar','A y D','Flechas izquierda y derecha','Horizontal1 o Horizontal2'),
('Cargar y disparar','Espacio','Enter','Fire1 o Fire2'),
('Pausar o reanudar','Escape o botón','Escape o botón','KeyCode Escape')],[1.55,1.45,1.8,1.75])
doc.add_heading('Movimiento',level=2)
code(doc,"""m_MovementAxisName = "Vertical" + m_PlayerNumber;
m_TurnAxisName = "Horizontal" + m_PlayerNumber;

m_MovementInputValue = Input.GetAxis(m_MovementAxisName);
m_TurnInputValue = Input.GetAxis(m_TurnAxisName);""",'El número del jugador selecciona sus ejes')
code(doc,"""Vector3 movement = transform.forward
    * m_MovementInputValue * m_Speed * Time.deltaTime;
m_Rigidbody.MovePosition(m_Rigidbody.position + movement);

float turn = m_TurnInputValue * m_TurnSpeed * Time.deltaTime;
Quaternion turnRotation = Quaternion.Euler(0f, turn, 0f);
m_Rigidbody.MoveRotation(m_Rigidbody.rotation * turnRotation);""",'La física recibe desplazamiento y rotación')
source(doc,'Assets/Scripts/Tank/TankMovement.cs','42 a 55 y 96 a 115')
page(doc)

doc.add_heading('6 Disparo cargado',level=1)
doc.add_paragraph('El disparo no usa una fuerza fija. Al pulsar, se reinicia la carga; mientras la tecla permanece presionada, aumenta la fuerza y se actualiza el slider. Al soltar o alcanzar el máximo, Fire crea el proyectil.')
code(doc,"""m_FireButton = "Fire" + m_PlayerNumber;
m_ChargeSpeed = (m_MaxLaunchForce - m_MinLaunchForce)
    / m_MaxChargeTime;

if (Input.GetButton(m_FireButton) && !m_Fired)
{
    m_CurrentLaunchForce += m_ChargeSpeed * Time.deltaTime;
    m_AimSlider.value = m_CurrentLaunchForce;
}""",'Carga gradual del disparo')
code(doc,"""Rigidbody shellInstance = Instantiate(
    m_Shell,
    m_FireTransform.position,
    m_FireTransform.rotation) as Rigidbody;

shellInstance.velocity =
    m_CurrentLaunchForce * m_FireTransform.forward;""",'Creación y lanzamiento del proyectil')
source(doc,'Assets/Scripts/Tank/TankShooting.cs','32 a 78 y 82 a 100')
doc.add_heading('Relación con el modelo de pato',level=2)
doc.add_paragraph('m_FireTransform es un objeto hijo colocado delante del personaje. En la versión de patitos se ajustó cerca del pico para que el proyectil parezca salir desde allí. La lógica de disparo no necesita conocer la forma del modelo; solo utiliza la posición y orientación de ese Transform.')
page(doc)

doc.add_heading('7 Personajes y apariencia',level=1)
doc.add_paragraph('Cada pato jugable nace del prefab Tank. El nombre técnico se conserva porque los scripts originales dependen de esa estructura. El modelo rubberDuck.fbx y el material DuckPlayer.mat sustituyen la apariencia, mientras que Rigidbody, Collider, Canvas, TankMovement, TankShooting y TankHealth mantienen el comportamiento.')
doc.add_heading('Configuración por jugador',level=2)
code(doc,"""m_Movement = m_Instance.GetComponent<TankMovement>();
m_Shooting = m_Instance.GetComponent<TankShooting>();

m_Movement.m_PlayerNumber = m_PlayerNumber;
m_Shooting.m_PlayerNumber = m_PlayerNumber;

MeshRenderer[] renderers =
    m_Instance.GetComponentsInChildren<MeshRenderer>();
for (int i = 0; i < renderers.Length; i++)
    renderers[i].material.color = m_PlayerColor;""",'TankManager conecta identidad, controles y color')
source(doc,'Assets/Scripts/Managers/TankManager.cs','25 a 47')
doc.add_heading('Activación durante las rondas',level=2)
code(doc,"""public void DisableControl()
{
    m_Movement.enabled = false;
    m_Shooting.enabled = false;
    m_CanvasGameObject.SetActive(false);
}

public void EnableControl()
{
    m_Movement.enabled = true;
    m_Shooting.enabled = true;
    m_CanvasGameObject.SetActive(true);
}""",'El administrador decide cuándo puede actuar cada personaje')
doc.add_paragraph('Esto evita que un jugador se mueva durante los mensajes de inicio y final de ronda. Reset devuelve cada instancia a su punto de aparición y reactiva sus componentes para recuperar la vida inicial.')
page(doc)

doc.add_heading('8 Vida impacto y eliminación',level=1)
code(doc,"""public void TakeDamage(float amount)
{
    m_CurrentHealth -= amount;
    SetHealthUI();

    if (m_CurrentHealth <= 0f && !m_Dead)
        OnDeath();
}""",'La salud centraliza daño y muerte')
code(doc,"""Vector3 explosionToTarget = targetPosition - transform.position;
float explosionDistance = explosionToTarget.magnitude;
float relativeDistance =
    (m_ExplosionRadius - explosionDistance) / m_ExplosionRadius;
float damage = relativeDistance * m_MaxDamage;
return Mathf.Max(0f, damage);""",'El daño disminuye con la distancia a la explosión')
source(doc,'Assets/Scripts/Tank/TankHealth.cs','37 a 67')
source(doc,'Assets/Scripts/Shell/ShellExplosion.cs','67 a 83')
doc.add_paragraph('Cuando la vida llega a cero, el pato se desactiva, se reproducen partículas y sonido, y GameManager detecta que queda uno o ningún personaje activo. Ese resultado termina la ronda y actualiza el marcador de victorias.')
doc.add_heading('Cadena de responsabilidades',level=2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('TankShooting  →  ShellExplosion  →  TankHealth  →  GameManager'); r.bold=True; r.font.size=Pt(11)
page(doc)

doc.add_heading('9 Ciclo de rondas',level=1)
code(doc,"""private IEnumerator GameLoop()
{
    yield return StartCoroutine(RoundStarting());
    yield return StartCoroutine(RoundPlaying());
    yield return StartCoroutine(RoundEnding());

    if (m_GameWinner != null)
        Application.LoadLevel(Application.loadedLevel);
    else
        StartCoroutine(GameLoop());
}""",'Secuencia principal de la partida')
source(doc,'Assets/Scripts/Managers/GameManager.cs','71 a 93')
table(doc,['Fase','Acción del código','Estado de los jugadores'],[
('RoundStarting','Reinicia posiciones, ajusta cámara, aumenta la ronda y espera.','Controles desactivados'),
('RoundPlaying','Limpia el mensaje y espera hasta que quede como máximo un personaje.','Controles activados'),
('RoundEnding','Calcula ganador, aumenta victorias y muestra resultados.','Controles desactivados')],[1.35,3.5,1.65])
doc.add_paragraph('Las corrutinas permiten repartir la secuencia entre varios cuadros sin bloquear Unity. WaitForSeconds respeta Time.timeScale, por lo que una pausa también congela las esperas activas de las rondas.')
page(doc)

doc.add_heading('10 Conexiones en la escena Main',level=1)
doc.add_paragraph('El código funciona porque la escena y los prefabs enlazan referencias públicas desde el Inspector. Si una referencia cambia al reemplazar assets, conviene revisar esta lista antes de modificar scripts.')
table(doc,['Objeto o componente','Referencias que deben estar asignadas'],[
('GameManager','CameraControl, MessageText, Tank prefab, lista de jugadores y SpawnPoints.'),
('GameHUD con GameTimerPause','TimerText, PauseButton y texto interno del botón.'),
('Prefab Tank con pato','Rigidbody, colliders, Canvas, FireTransform, AudioSources y scripts Tank.'),
('TankShooting','Shell prefab, FireTransform, AimSlider y clips de carga y disparo.'),
('TankHealth','Slider, FillImage y prefab de explosión.'),
('Shell','Capa Players, partículas, AudioSource, radio y daño máximo.')],[2.25,4.25])
doc.add_heading('Comprobación rápida',level=2)
bullet(doc,'Main aparece como escena principal y única escena necesaria en Build Settings.')
bullet(doc,'El temporizador empieza en 00:00 y se detiene al pausar.')
bullet(doc,'El botón alterna entre PAUSAR y REANUDAR; Escape produce el mismo resultado.')
bullet(doc,'Los dos patitos reciben colores diferentes y responden a sus respectivos controles.')
bullet(doc,'El proyectil nace junto al pico, causa daño y puede finalizar una ronda.')
page(doc)

doc.add_heading('11 Mejoras recomendadas',level=1)
doc.add_paragraph('Las siguientes mejoras son coherentes con la estructura actual y pueden presentarse como trabajo futuro. No forman parte del código implementado salvo que se indique lo contrario.')
table(doc,['Prioridad','Mejora','Cambio técnico sugerido'],[
('Alta','Menú de inicio real','Crear MenuController y comenzar GameLoop mediante un método público después del clic.'),
('Alta','Reinicio desde pausa','Añadir botón Reiniciar que restaure Time.timeScale antes de recargar Main.'),
('Media','Tiempo por ronda','Reiniciar elapsedTime en RoundStarting o mantener contadores separados de partida y ronda.'),
('Media','Nuevo Input System','Sustituir nombres de ejes por acciones y mapas para teclado y mandos.'),
('Media','Nombres temáticos','Renombrar clases gradualmente a Duck sin romper referencias serializadas del prefab.'),
('Baja','Indicador de estado','Mostrar PAUSA en el centro de la pantalla además de cambiar el texto del botón.')],[.75,1.65,4.1])
page(doc)
doc.add_heading('Archivos principales para la exposición',level=2)
for x in ['Assets/Scripts/Managers/GameTimerPause.cs','Assets/Scripts/Managers/GameManager.cs','Assets/Scripts/Managers/TankManager.cs','Assets/Scripts/Tank/TankMovement.cs','Assets/Scripts/Tank/TankShooting.cs','Assets/Scripts/Tank/TankHealth.cs','Assets/Scripts/Shell/ShellExplosion.cs']:
    bullet(doc,x)
doc.add_heading('Cierre',level=2)
doc.add_paragraph('El proyecto separa correctamente la lógica general, la identidad de cada jugador y los comportamientos del personaje. La mejora de temporizador y pausa se integró sin alterar el sistema de rondas. Esta separación permite cambiar la apariencia del tanque por un pato y conservar las mecánicas centrales del juego.')

doc.core_properties.title='Guía técnica del código Patitos en batalla'
doc.core_properties.subject='Explicación del temporizador, pausa, inicio, controles y personajes del proyecto Unity'
doc.core_properties.author='Proyecto Patitos en batalla'
doc.save(OUT)
print(OUT)
