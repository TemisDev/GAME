from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(r"C:\Users\vince\OneDrive\Documents\ChatGPT\Game development\GDD Patitos en batalla.docx")
DUCK = Path(r"C:\Users\vince\OneDrive\Documents\ChatGPT\Game development\DuckSwap\DuckPreview.png")

NAVY = "18232E"
ORANGE = "F6A623"
YELLOW = "FFD83D"
PALE = "FFF7D6"
LIGHT = "F3F5F7"
MID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_table(table, color=MID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_runs(paragraph, size=10.5, bold=False, color=BLACK):
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    format_runs(p)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)
        format_runs(p)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    border_table(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_width(cell, widths[i])
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_runs(cell.paragraphs[0], 9.5, True, RGBColor(255, 255, 255))
        set_cell_margins(cell)
    for ri, data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, text in enumerate(data):
            cell = row.cells[i]
            set_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                shade(cell, LIGHT)
            cell.text = str(text)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
            format_runs(cell.paragraphs[0], 9.2)
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = BLACK
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(30)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = BLACK
for name, size in (("Heading 1", 20), ("Heading 2", 14)):
    styles[name].font.name = "Aptos Display"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = BLACK
    styles[name].paragraph_format.space_before = Pt(10)
    styles[name].paragraph_format.space_after = Pt(6)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Documento de diseño del videojuego")
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Patitos en batalla")
r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(25); r.font.color.rgb = BLACK
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p3.add_run("Juego de combate local para dos jugadores desarrollado en Unity")
r.font.size = Pt(12); r.font.color.rgb = RGBColor(70, 70, 70)
doc.add_paragraph()
if DUCK.exists():
    ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(str(DUCK), width=Inches(3.6))
doc.add_paragraph()
meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER; meta.autofit = False; border_table(meta)
for i, (k, v) in enumerate([
    ("Motor", "Unity 2019.4.19f1"),
    ("Plataforma", "Navegador WebGL y equipo de escritorio"),
    ("Versión documentada", "Escena Main con patitos, música nueva, cronómetro y pausa")]):
    set_width(meta.rows[i].cells[0], 1.7); set_width(meta.rows[i].cells[1], 4.5)
    meta.rows[i].cells[0].text = k; meta.rows[i].cells[1].text = v
    shade(meta.rows[i].cells[0], PALE)
    format_runs(meta.rows[i].cells[0].paragraphs[0], 9.5, True)
    format_runs(meta.rows[i].cells[1].paragraphs[0], 9.5)
    for c in meta.rows[i].cells: set_cell_margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

page_break(doc)
add_heading(doc, "1 Resumen del proyecto", 1)
add_body(doc, "Patitos en batalla es un videojuego de combate local para dos participantes que transforma la mecánica clásica de Tanks en una experiencia humorística. Cada jugador controla un patito de goma, se desplaza por un escenario desértico y lanza proyectiles para ganar rondas. La propuesta conserva la claridad del juego original y cambia su identidad visual y sonora para producir una experiencia más ligera y memorable.")
add_heading(doc, "Objetivo del documento", 2)
add_body(doc, "Este GDD define la visión, las reglas, los personajes, los recursos audiovisuales, los controles, la interfaz y la estructura técnica de la versión actual. También registra las mejoras incorporadas para que el proyecto pueda explicarse, probarse y ampliarse de manera coherente.")
add_heading(doc, "Ficha general", 2)
add_table(doc, ["Elemento", "Definición"], [
    ("Género", "Acción competitiva y disparos con vista cenital"),
    ("Modo", "Multijugador local para dos personas en el mismo teclado"),
    ("Público", "Jugadores casuales, estudiantes y público familiar"),
    ("Duración", "Partidas breves organizadas en rondas"),
    ("Cámara", "Cámara dinámica que mantiene visibles a ambos jugadores"),
    ("Escena principal", "Assets/Scenes/Main.unity"),
], [1.65, 4.95])
add_heading(doc, "Pilares de diseño", 2)
add_bullets(doc, [
    "Competencia inmediata: controles simples y objetivo reconocible desde la primera ronda.",
    "Humor visual: los vehículos militares se sustituyen por patitos de goma de colores.",
    "Lectura clara: cámara compartida, indicadores de vida, mensajes de ronda y temporizador visible.",
    "Sesiones cortas: reinicio rápido de rondas y puntuación acumulada hasta definir al ganador.",
])

page_break(doc)
add_heading(doc, "2 Concepto y experiencia del jugador", 1)
add_heading(doc, "Premisa", 2)
add_body(doc, "Dos patitos de goma se enfrentan en una arena desértica repleta de edificios, rocas y estructuras industriales. Cada participante debe moverse, orientar su personaje, cargar el disparo y calcular la trayectoria del proyectil. El último patito activo gana la ronda.")
add_heading(doc, "Objetivo principal", 2)
add_body(doc, "Eliminar al oponente antes de perder toda la vida. Las victorias se acumulan durante varias rondas y el primer participante que alcanza el número de triunfos configurado gana la partida completa.")
add_heading(doc, "Bucle principal de juego", 2)
add_table(doc, ["Etapa", "Acción del sistema", "Decisión del jugador"], [
    ("Inicio", "Reposiciona personajes y ajusta la cámara", "Reconocer la posición inicial"),
    ("Combate", "Habilita movimiento, disparo y daño", "Moverse, apuntar y cargar el disparo"),
    ("Resolución", "Detecta al último personaje activo", "Observar el resultado de la ronda"),
    ("Puntuación", "Suma la victoria y muestra el marcador", "Prepararse para la ronda siguiente"),
    ("Fin", "Declara al ganador y reinicia la escena", "Comenzar una nueva partida"),
], [1.15, 2.55, 2.9])
add_heading(doc, "Condiciones", 2)
add_bullets(doc, [
    "Victoria de ronda: solamente un patito permanece activo.",
    "Empate: ambos personajes quedan eliminados.",
    "Victoria de partida: un jugador alcanza cinco rondas ganadas según la configuración actual.",
    "Pausa: el botón o la tecla Escape detienen el tiempo y todas las acciones dependientes del juego.",
])

page_break(doc)
add_heading(doc, "3 Personajes jugables", 1)
add_body(doc, "Los dos personajes comparten las mismas estadísticas y capacidades. La diferencia principal es el color asignado a cada jugador, lo cual permite identificarlos rápidamente sin generar ventajas competitivas.")
add_table(doc, ["Personaje", "Identidad visual", "Función", "Capacidades"], [
    ("Patito jugador 1", "Color configurado para el jugador 1", "Combatiente local", "Avanzar, retroceder, girar, cargar y disparar"),
    ("Patito jugador 2", "Color configurado para el jugador 2", "Combatiente local", "Avanzar, retroceder, girar, cargar y disparar"),
], [1.3, 1.75, 1.45, 2.1])
add_heading(doc, "Composición del personaje", 2)
add_bullets(doc, [
    "Modelo visual: Rubber Duck de J Toastie, importado en formato FBX.",
    "Material: textura original con adaptación del color corporal para distinguir jugadores.",
    "Colisión y física: se conserva el cuerpo físico del tanque original para mantener el comportamiento estable.",
    "Vida: barra de salud en espacio de mundo asociada a cada personaje.",
    "Disparo: el proyectil aparece delante del pico del patito.",
    "Audio: efectos de carga, disparo, impacto y eliminación heredados del sistema original.",
])
add_heading(doc, "Balance", 2)
add_body(doc, "Ambos jugadores utilizan la misma velocidad, fuerza máxima de disparo, resistencia y tamaño de colisión. Esta simetría concentra la competencia en el posicionamiento, el tiempo de carga y la precisión.")

page_break(doc)
add_heading(doc, "4 Controles", 1)
add_body(doc, "El juego utiliza un solo teclado. Los controles están separados para evitar conflictos y permitir que dos personas jueguen simultáneamente.")
add_table(doc, ["Acción", "Jugador 1", "Jugador 2"], [
    ("Avanzar", "W", "Flecha arriba"),
    ("Retroceder", "S", "Flecha abajo"),
    ("Girar a la izquierda", "A", "Flecha izquierda"),
    ("Girar a la derecha", "D", "Flecha derecha"),
    ("Cargar y disparar", "Espacio", "Enter"),
    ("Pausar o reanudar", "Escape o botón en pantalla", "Escape o botón en pantalla"),
], [2.25, 2.1, 2.25])
add_heading(doc, "Comportamiento del disparo", 2)
add_body(doc, "Al mantener pulsado el botón de disparo aumenta la fuerza del lanzamiento. Al soltarlo, el proyectil sale desde el pico en la dirección del personaje. Si se alcanza la fuerza máxima, el disparo se ejecuta automáticamente.")
add_heading(doc, "Principios de accesibilidad", 2)
add_bullets(doc, [
    "Asignaciones cercanas para cada jugador y sin combinaciones complejas.",
    "Colores contrastantes para reconocer a los personajes.",
    "Mensajes grandes para rondas, resultados y ganador final.",
    "Botón visible para la pausa, además del atajo de teclado.",
])

page_break(doc)
add_heading(doc, "5 Escenario y assets", 1)
add_heading(doc, "Escena principal", 2)
add_body(doc, "La versión jugable utiliza Assets/Scenes/Main.unity como única escena habilitada en Build Settings. En ella se encuentran el escenario, la cámara, los puntos de aparición, el sistema de eventos, la interfaz, el GameManager y el nuevo GameHUD.")
add_heading(doc, "Entorno", 2)
add_body(doc, "La arena presenta una estética low poly desértica. Las estructuras generan rutas, espacios de cobertura y obstáculos que alteran el movimiento y la trayectoria de los proyectiles.")
add_table(doc, ["Categoría", "Assets principales", "Uso"], [
    ("Terreno", "Terrain, Dunes01, Dunes02, Crater01", "Base visual y relieve de la arena"),
    ("Edificaciones", "Building01, Building02, Refinery, OilStorage", "Cobertura y puntos de referencia"),
    ("Industriales", "PumpJack, Radar, Helipad", "Ambientación y obstáculos"),
    ("Naturaleza", "Rocks01 a Rocks03, Cliff, Cactus, PalmTree, Tree", "Límites visuales y variedad"),
    ("Personajes", "rubberDuck.fbx y DuckPlayer.mat", "Representación de ambos jugadores"),
    ("Combate", "Shell, ShellExplosion y TankExplosion", "Proyectiles, impactos y eliminación"),
], [1.2, 2.65, 2.75])
add_heading(doc, "Dirección artística", 2)
add_body(doc, "La mezcla de un escenario militar y personajes de juguete crea el contraste humorístico central. Las formas simples, los colores vivos y la cámara cenital mantienen la legibilidad aun cuando los personajes se separan.")

page_break(doc)
add_heading(doc, "6 Sistemas de juego", 1)
add_table(doc, ["Sistema", "Responsabilidad", "Implementación principal"], [
    ("Movimiento", "Traslación y giro mediante física", "TankMovement"),
    ("Disparo", "Carga, lanzamiento y audio del proyectil", "TankShooting"),
    ("Salud", "Daño, barra de vida y eliminación", "TankHealth"),
    ("Explosión", "Área de daño y efectos al impactar", "ShellExplosion"),
    ("Rondas", "Inicio, combate, resultado y reinicio", "GameManager"),
    ("Jugadores", "Colores, entradas, control y victorias", "TankManager"),
    ("Cámara", "Encuadre dinámico de ambos personajes", "CameraControl"),
    ("HUD", "Cronómetro y pausa", "GameTimerPause"),
], [1.25, 2.9, 2.45])
add_heading(doc, "Cámara dinámica", 2)
add_body(doc, "La cámara calcula una posición media entre los jugadores y ajusta el zoom para mantenerlos visibles. Al comenzar cada ronda vuelve inmediatamente al encuadre adecuado.")
add_heading(doc, "Gestión de rondas", 2)
add_body(doc, "GameManager instancia los personajes, bloquea los controles durante las transiciones, habilita el combate, detecta el final de la ronda y actualiza el marcador. Cuando un participante consigue cinco victorias, la escena se recarga para iniciar una partida nueva.")

page_break(doc)
add_heading(doc, "7 Interfaz y mejoras incorporadas", 1)
add_heading(doc, "Cronómetro", 2)
add_body(doc, "El GameHUD muestra el tiempo transcurrido en formato minutos y segundos. El conteo comienza en 00:00, continúa durante las rondas y se detiene cuando el juego está pausado. Esta información ayuda a comparar la duración de partidas y aporta una referencia constante de progreso.")
add_heading(doc, "Botón de pausa", 2)
add_body(doc, "El botón ubicado en la parte superior derecha alterna entre PAUSAR y REANUDAR. Al activarlo, Time.timeScale pasa a cero, por lo que se detienen el movimiento, los proyectiles, las transiciones y el cronómetro. La tecla Escape ejecuta la misma acción.")
add_heading(doc, "Mensajes de partida", 2)
add_bullets(doc, [
    "Título inicial y mensajes centrales mediante MessageCanvas.",
    "Número de ronda antes de habilitar el combate.",
    "Ganador de la ronda y puntuación acumulada.",
    "Ganador final cuando alcanza el límite configurado.",
    "Barras de salud vinculadas a cada personaje.",
])
add_heading(doc, "Jerarquía relevante de Main", 2)
add_table(doc, ["Objeto", "Función"], [
    ("LevelArt", "Contenido visual y colisiones del escenario"),
    ("CameraRig", "Seguimiento y encuadre de los jugadores"),
    ("SpawnPoint1 y SpawnPoint2", "Posiciones iniciales de los participantes"),
    ("MessageCanvas", "Mensajes de ronda y resultados"),
    ("GameManager", "Control general de partida"),
    ("GameHUD", "Cronómetro y botón de pausa"),
], [2.4, 4.2])

page_break(doc)
add_heading(doc, "8 Audio", 1)
add_body(doc, "El diseño sonoro separa la música ambiental de las respuestas del combate. Los efectos comunican el estado de carga, la ejecución del disparo, el impacto y la eliminación sin depender únicamente de señales visuales.")
add_table(doc, ["Elemento", "Recurso", "Función"], [
    ("Música ambiental", "the_mountain-game-game-music-508018.mp3", "Acompañamiento continuo durante la partida"),
    ("Movimiento", "EngineIdle y EngineDriving", "Respuesta al reposo y desplazamiento"),
    ("Carga", "ShotCharging.wav", "Indica el aumento de fuerza"),
    ("Disparo", "ShotFiring.wav", "Confirma el lanzamiento"),
    ("Impacto", "ShellExplosion.wav", "Refuerza el daño de área"),
    ("Eliminación", "TankExplosion.wav", "Comunica la derrota del personaje"),
], [1.35, 2.8, 2.45])
add_heading(doc, "Criterios de mezcla", 2)
add_bullets(doc, [
    "La música debe permanecer por debajo de los efectos de disparo e impacto.",
    "El bucle ambiental debe evitar cortes perceptibles.",
    "Los efectos deben ser breves para no saturar rondas rápidas.",
    "El tono sonoro puede evolucionar hacia chirridos de juguete, burbujas y cuacs si se desea reforzar la identidad cómica.",
])

page_break(doc)
add_heading(doc, "9 Arquitectura técnica", 1)
add_table(doc, ["Elemento", "Decisión técnica"], [
    ("Motor", "Unity 2019.4.19f1"),
    ("Lenguaje", "C sharp"),
    ("Física", "Rigidbody y sistema de colisiones tridimensional"),
    ("Interfaz", "Unity UI con Canvas, Text, Button y EventSystem"),
    ("Entrada", "Input Manager clásico con ejes separados por jugador"),
    ("Publicación", "Compilación WebGL para Unity Play"),
    ("Escena de arranque", "Assets/Scenes/Main.unity en el índice 0"),
], [2.0, 4.6])
add_heading(doc, "Organización del proyecto", 2)
add_bullets(doc, [
    "Assets/Scenes contiene la escena principal Main.",
    "Assets/Prefabs contiene personajes, proyectiles, explosiones y escenario reutilizable.",
    "Assets/Scripts agrupa cámara, administradores, interfaz, proyectiles y lógica de personajes.",
    "Assets/AudioClips contiene música y efectos.",
    "Assets/Ducks contiene el modelo, textura, material y crédito del patito.",
    "Assets/Models contiene los modelos originales del entorno low poly.",
])
add_heading(doc, "Persistencia de cambios", 2)
add_body(doc, "El cronómetro y el botón están guardados como objetos de Main, y la música está referenciada desde las escenas. Para publicar una versión nueva es necesario generar otra compilación WebGL y comprimir su contenido en ZIP.")

page_break(doc)
add_heading(doc, "10 Alcance y mejoras futuras", 1)
add_heading(doc, "Estado actual", 2)
add_table(doc, ["Función", "Estado"], [
    ("Combate local para dos jugadores", "Implementado"),
    ("Patitos diferenciados por color", "Implementado"),
    ("Escena Main como inicio", "Implementado"),
    ("Música ambiental personalizada", "Implementado"),
    ("Cronómetro de partida", "Implementado"),
    ("Pausa mediante botón y Escape", "Implementado"),
    ("Publicación WebGL", "Requiere recompilar después de los últimos cambios"),
], [3.7, 2.9])
add_heading(doc, "Mejoras recomendadas", 2)
add_table(doc, ["Prioridad", "Mejora", "Beneficio"], [
    ("Alta", "Pantalla inicial con instrucciones", "Reduce dudas sobre controles y objetivo"),
    ("Alta", "Recompilar y probar WebGL", "Garantiza que el HUD y la música estén en Unity Play"),
    ("Media", "Proyectiles con apariencia de burbujas", "Refuerza la temática de patitos"),
    ("Media", "Efectos de cuac y salpicadura", "Crea una identidad sonora propia"),
    ("Media", "Selector de número de rondas", "Permite ajustar la duración de la partida"),
    ("Baja", "Compatibilidad con controles", "Amplía comodidad y accesibilidad"),
    ("Baja", "Modo contra inteligencia artificial", "Permite jugar sin un segundo participante"),
], [1.0, 2.7, 2.9])
add_heading(doc, "Riesgos de producción", 2)
add_bullets(doc, [
    "Los dos jugadores comparten teclado; algunos equipos limitan ciertas combinaciones simultáneas.",
    "Las modificaciones posteriores requieren volver a generar el paquete WebGL.",
    "La música y los modelos externos deben conservar su información de licencia y atribución.",
    "Los cambios de escala visual no deben alterar colisiones ni equilibrio entre personajes.",
])

page_break(doc)
add_heading(doc, "11 Créditos y referencias de recursos", 1)
add_body(doc, "El proyecto parte del tutorial Tanks de Unity y conserva su arquitectura de cámara, combate, rondas, interfaz y recursos de escenario. La adaptación transforma los vehículos en patitos, incorpora una identidad sonora diferente y añade funciones de interfaz.")
add_table(doc, ["Recurso", "Procedencia o autor", "Uso en el proyecto"], [
    ("Base del juego", "Unity Tanks Tutorial", "Mecánicas, scripts y assets del escenario"),
    ("Modelo Rubber Duck", "J Toastie en Poly Pizza", "Personajes jugables; licencia Creative Commons Attribution"),
    ("Música ambiental", "Archivo descargado the_mountain-game-game-music-508018.mp3", "Fondo musical; conservar la licencia de la página de descarga"),
    ("Adaptación", "Proyecto Patitos en batalla", "Colores, sustitución visual, cronómetro, pausa y configuración de Main"),
], [1.35, 2.25, 3.0])
add_heading(doc, "Conclusión", 2)
add_body(doc, "Patitos en batalla demuestra cómo una base jugable conocida puede adquirir una identidad distinta mediante cambios de personajes, audio e interfaz. La versión actual ofrece una partida completa para dos jugadores, mantiene reglas fáciles de aprender e incorpora herramientas de calidad de vida que facilitan el control del tiempo y la interrupción segura de la sesión.")

# Footer
for sec in doc.sections:
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Patitos en batalla  |  Documento de diseño del videojuego")
    run.font.name = "Aptos"; run.font.size = Pt(8); run.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "Documento de diseño del videojuego Patitos en batalla"
doc.core_properties.subject = "Game Design Document"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
